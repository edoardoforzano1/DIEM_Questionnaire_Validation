from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from pathlib import Path
from shutil import copy2

import openpyxl
from openpyxl.styles import PatternFill

from admin_tools_common import (
    as_bool,
    detect_label_column,
    merge_runtime_config,
    now_stamp,
    print_runtime_banner,
    resolve_source_workbook,
)


FILL_REMAP = PatternFill("solid", fgColor="D9EAD3")


def _norm(value) -> str:
    return str(value or "").strip().lower()


def _get_choices_context(wb, language: str):
    if "choices" not in wb.sheetnames:
        raise KeyError("Workbook has no 'choices' sheet.")

    ws = wb["choices"]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise ValueError("choices sheet is empty.")

    headers = [str(v).strip() if v is not None else "" for v in header_row]
    idx = {h: i for i, h in enumerate(headers)}
    list_col = idx.get("list_name")
    name_col = idx.get("name")
    if list_col is None or name_col is None:
        raise ValueError("choices sheet must contain 'list_name' and 'name' columns.")

    label_col = detect_label_column(headers, language)
    parent_col = None
    parent_header = ""
    for candidate in ("my_filter_admin", "choice_filter", "filter", "tema", "team"):
        if candidate in idx:
            parent_col = idx[candidate]
            parent_header = candidate
            break
    if parent_col is None:
        raise ValueError("choices sheet must contain an admin parent column such as 'my_filter_admin'.")

    return ws, headers, list_col, name_col, label_col, parent_col, parent_header


def _collect_admin_rows(ws, list_col: int, name_col: int, label_col: int | None, parent_col: int):
    admin2_rows: list[dict] = []
    admin3_rows: list[dict] = []

    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        list_name = str((row[list_col] if list_col < len(row) else "") or "").strip().lower()
        if list_name not in {"admin2", "admin3"}:
            continue

        code = str((row[name_col] if name_col < len(row) else "") or "").strip()
        if not code:
            continue

        label = str((row[label_col] if (label_col is not None and label_col < len(row)) else "") or "").strip()
        parent = str((row[parent_col] if parent_col < len(row) else "") or "").strip()
        item = {
            "excel_row": excel_row,
            "list_name": list_name,
            "code": code,
            "label": label,
            "parent": parent,
        }
        if list_name == "admin2":
            admin2_rows.append(item)
        else:
            admin3_rows.append(item)

    return admin2_rows, admin3_rows


def _expand_merged_admin_code(merged_code: str) -> list[str]:
    parts = [str(p).strip() for p in str(merged_code or "").split("_") if str(p).strip()]
    if len(parts) < 2:
        return []

    base = parts[0]
    if len(base) < 4 or not base[-3:].isdigit():
        return []
    prefix, tail = base[:-3], base[-3:]
    if any(not p.isdigit() for p in parts[1:]):
        return []

    width = len(tail)
    components = [base]
    for suffix in parts[1:]:
        components.append(prefix + suffix.zfill(width))
    return components


def _infer_admin2_merge_map(admin2_rows: list[dict]):
    reverse_map: dict[str, set[str]] = defaultdict(set)
    merge_defs: list[dict] = []
    admin2_by_code = {
        str(row.get("code") or "").strip(): row
        for row in admin2_rows
        if str(row.get("code") or "").strip()
    }

    for row in admin2_rows:
        merged_code = str(row.get("code") or "").strip()
        components = _expand_merged_admin_code(merged_code)
        if not components:
            continue
        merge_defs.append(
            {
                "merged_code": merged_code,
                "merged_label": str(row.get("label") or "").strip(),
                "parent": str(row.get("parent") or "").strip(),
                "components": components,
            }
        )
        for component in components:
            reverse_map[component].add(merged_code)

    return admin2_by_code, reverse_map, merge_defs


def _analyse_admin3_parent_rows(admin2_rows: list[dict], admin3_rows: list[dict]):
    admin2_by_code, reverse_map, merge_defs = _infer_admin2_merge_map(admin2_rows)
    admin2_codes = set(admin2_by_code)
    details: list[dict] = []

    for row in admin3_rows:
        parent = str(row.get("parent") or "").strip()
        if not parent:
            details.append(
                {
                    **row,
                    "action": "missing_parent",
                    "new_parent": "",
                    "new_parent_label": "",
                    "candidate_merged_codes": "",
                    "problem": "admin3 row has blank parent code",
                }
            )
            continue

        if parent in admin2_codes:
            continue

        candidates = sorted(reverse_map.get(parent, set()))
        if len(candidates) == 1:
            merged_code = candidates[0]
            merged_label = str((admin2_by_code.get(merged_code) or {}).get("label") or "").strip()
            details.append(
                {
                    **row,
                    "action": "replace",
                    "new_parent": merged_code,
                    "new_parent_label": merged_label,
                    "candidate_merged_codes": merged_code,
                    "problem": "invalid parent remapped to inferred merged admin2 code",
                }
            )
        elif len(candidates) > 1:
            details.append(
                {
                    **row,
                    "action": "ambiguous",
                    "new_parent": "",
                    "new_parent_label": "",
                    "candidate_merged_codes": ", ".join(candidates),
                    "problem": "invalid parent matches more than one merged admin2 code",
                }
            )
        else:
            details.append(
                {
                    **row,
                    "action": "unresolved",
                    "new_parent": "",
                    "new_parent_label": "",
                    "candidate_merged_codes": "",
                    "problem": "invalid parent not found in inferred merged admin2 definitions",
                }
            )

    return details, merge_defs


def _apply_parent_updates(ws, parent_col: int, details: list[dict]) -> int:
    changed = 0
    for row in details:
        if row.get("action") != "replace":
            continue
        excel_row = int(row.get("excel_row") or 0)
        if excel_row < 2:
            continue
        cell = ws.cell(row=excel_row, column=parent_col + 1)
        cell.value = row.get("new_parent") or ""
        cell.fill = FILL_REMAP
        changed += 1
    return changed


def _write_detail_csv(path: Path, details: list[dict]) -> Path:
    fieldnames = [
        "action",
        "excel_row",
        "admin3_code",
        "admin3_label",
        "old_parent",
        "new_parent",
        "new_parent_label",
        "candidate_merged_codes",
        "problem",
    ]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in details:
            writer.writerow(
                {
                    "action": row.get("action", ""),
                    "excel_row": row.get("excel_row", ""),
                    "admin3_code": row.get("code", ""),
                    "admin3_label": row.get("label", ""),
                    "old_parent": row.get("parent", ""),
                    "new_parent": row.get("new_parent", ""),
                    "new_parent_label": row.get("new_parent_label", ""),
                    "candidate_merged_codes": row.get("candidate_merged_codes", ""),
                    "problem": row.get("problem", ""),
                }
            )
    return path


def _write_summary_txt(
    path: Path,
    *,
    source: Path,
    output_file: Path,
    parent_header: str,
    merge_defs: list[dict],
    details: list[dict],
    changed_rows: int,
) -> Path:
    action_counts = Counter(str(row.get("action") or "") for row in details)
    lines = [
        "Admin3 Parent Remap Report",
        f"Source workbook : {source}",
        f"Output workbook : {output_file}",
        f"Parent column   : {parent_header}",
        f"Merged admin2 definitions detected : {len(merge_defs)}",
        f"Invalid admin3 parent rows scanned : {len(details)}",
        f"Rows remapped                    : {action_counts.get('replace', 0)}",
        f"Rows unresolved                 : {action_counts.get('unresolved', 0)}",
        f"Rows ambiguous                  : {action_counts.get('ambiguous', 0)}",
        f"Rows with blank parent          : {action_counts.get('missing_parent', 0)}",
        f"Edited cells highlighted        : {changed_rows}",
        "",
        "Merged admin2 definitions",
        "-------------------------",
    ]

    if not merge_defs:
        lines.append("(none detected)")
    else:
        for item in sorted(merge_defs, key=lambda x: str(x.get("merged_code") or "")):
            merged_code = str(item.get("merged_code") or "")
            merged_label = str(item.get("merged_label") or "")
            components = ", ".join(item.get("components") or [])
            lines.append(f"{merged_code} | {merged_label} -> {components}")

    lines.extend([
        "",
        "Invalid admin3 parent rows",
        "--------------------------",
    ])
    if not details:
        lines.append("(none)")
    else:
        for row in sorted(details, key=lambda x: int(x.get("excel_row") or 0)):
            lines.append(
                f"row {row.get('excel_row')} | {row.get('code')} | {row.get('label')} | "
                f"{row.get('parent')} -> {row.get('new_parent') or '(no change)'} | "
                f"{row.get('action')} | {row.get('problem')}"
            )

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def run(cfg: dict, *, dry_run_override: bool | None = None, output_dir_override: str = "") -> int:
    print_runtime_banner("tool_admin_parent_remap", cfg)
    source = resolve_source_workbook(cfg)
    print(f"  source      : {source}")

    dry_run = as_bool(dry_run_override, cfg.get("dry_run")) if dry_run_override is not None else as_bool(cfg.get("dry_run"), False)
    out_dir = Path(output_dir_override) if output_dir_override else Path(str(cfg.get("output_dir")))
    out_dir.mkdir(parents=True, exist_ok=True)

    stamp = now_stamp()
    out_file = out_dir / f"{source.stem}_admin3parentremap_{stamp}.xlsx"
    txt_file = out_dir / f"{source.stem}_admin3parentremap_{stamp}.txt"
    csv_file = out_dir / f"{source.stem}_admin3parentremap_details_{stamp}.csv"

    wb = openpyxl.load_workbook(source)
    ws, _, list_col, name_col, label_col, parent_col, parent_header = _get_choices_context(
        wb,
        str(cfg.get("language") or "en"),
    )
    admin2_rows, admin3_rows = _collect_admin_rows(ws, list_col, name_col, label_col, parent_col)
    details, merge_defs = _analyse_admin3_parent_rows(admin2_rows, admin3_rows)

    changed_rows = 0
    if dry_run:
        print("  dry_run     : true (no files written)")
    else:
        copy2(source, out_file)
        out_wb = openpyxl.load_workbook(out_file)
        out_ws, _, _, _, _, out_parent_col, _ = _get_choices_context(out_wb, str(cfg.get("language") or "en"))
        changed_rows = _apply_parent_updates(out_ws, out_parent_col, details)
        out_wb.save(out_file)
        _write_detail_csv(csv_file, details)
        _write_summary_txt(
            txt_file,
            source=source,
            output_file=out_file,
            parent_header=parent_header,
            merge_defs=merge_defs,
            details=details,
            changed_rows=changed_rows,
        )
        print(f"  output      : {out_file}")
        print(f"  report txt  : {txt_file}")
        print(f"  report csv  : {csv_file}")

    action_counts = Counter(str(row.get("action") or "") for row in details)
    print(
        "  remap stats : "
        + f"merged_admin2={len(merge_defs)} "
        + f"invalid_admin3={len(details)} "
        + f"replace={action_counts.get('replace', 0)} "
        + f"unresolved={action_counts.get('unresolved', 0)} "
        + f"ambiguous={action_counts.get('ambiguous', 0)} "
        + f"missing_parent={action_counts.get('missing_parent', 0)}"
    )
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Repair admin3 parent codes by remapping them to inferred merged admin2 codes."
    )
    parser.add_argument("--config", default="", help="Optional path to admin_tools_config.yaml")
    parser.add_argument("--table", dest="config", help=argparse.SUPPRESS)
    parser.add_argument("--source-file", default="", help="Optional explicit source workbook path override")
    parser.add_argument("--output-dir", default="", help="Optional explicit output folder override")
    parser.add_argument("--dry-run", action="store_true", help="Analyse only; do not write workbook or reports")
    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parents[1]
    cfg_path = Path(args.config).expanduser().resolve() if args.config else None
    cfg = merge_runtime_config(repo_root, explicit_cfg=cfg_path)
    if args.source_file:
        cfg["source_mode"] = "custom"
        cfg["source_file"] = args.source_file

    try:
        return run(
            cfg,
            dry_run_override=True if args.dry_run else None,
            output_dir_override=args.output_dir,
        )
    except Exception as e:
        print(f"ERROR: {e}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
